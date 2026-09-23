"""Utilidades para formateo de archivos Excel.

Proporciona funciones auxiliares para aplicar formato visual
a hojas de cálculo Excel generadas con openpyxl.
"""

from openpyxl.styles import PatternFill


def aplicar_color_encabezado(ws, fill: PatternFill) -> None:
    """Aplica color de fondo a todas las celdas de la primera fila (encabezados).

    Args:
        ws: Hoja de cálculo de openpyxl.
        fill: Objeto PatternFill con el color a aplicar.
    """
    for cell in ws[1]:
        cell.fill = fill
